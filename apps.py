import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Cm
import os, copy, tempfile, zipfile, traceback, re

# ================== CONFIG ==================
st.set_page_config(page_title="Tool Điền Đơn Online", layout="centered")

IMG_W = Cm(3)   # 3cm
IMG_H = Cm(4)   # 4cm

# ================== TEXT UTILS ==================
def get_all_text(doc):
    texts = []

    for p in doc.paragraphs:
        texts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    texts.append(p.text)

    return texts


def extract_placeholders(doc):
    texts = get_all_text(doc)
    placeholders = set()

    for t in texts:
        found = re.findall(r"\{\{(.*?)\}\}", t)
        for f in found:
            placeholders.add(f.strip())

    return list(placeholders)

def replace_text_in_paragraph(paragraph, placeholder, new_text):
    full_text = "".join(run.text for run in paragraph.runs)

    if placeholder in full_text:
        full_text = full_text.replace(placeholder, new_text)

        for run in paragraph.runs:
            run.text = ""

        paragraph.runs[0].text = full_text


def replace_text_in_cell(cell, placeholder, new_text):
    for p in cell.paragraphs:
        replace_text_in_paragraph(p, placeholder, new_text)

# ================== DATE ==================
def normalize_date(value):
    try:
        dt = pd.to_datetime(value, dayfirst=True)
        return dt.strftime("%d/%m/%Y")
    except:
        return str(value)

# ================== IMAGE ==================
def insert_image(paragraph_or_cell, placeholder, image_path):
    if not os.path.isfile(image_path):
        return

    if hasattr(paragraph_or_cell, "paragraphs"):
        paragraphs = paragraph_or_cell.paragraphs
    else:
        paragraphs = [paragraph_or_cell]

    for p in paragraphs:
        full_text = "".join(run.text for run in p.runs)

        if placeholder in full_text:
            # Xóa text placeholder
            full_text = full_text.replace(placeholder, "")

            for run in p.runs:
                run.text = ""

            p.runs[0].text = full_text

            # Thêm ảnh
            new_run = p.add_run()
            new_run.add_picture(image_path, width=IMG_W, height=IMG_H)

# ================== ZIP IMAGE ==================
def extract_images(zip_file, temp_dir):
    image_map = {}

    with zipfile.ZipFile(zip_file, 'r') as z:
        z.extractall(temp_dir)

    for root, _, files in os.walk(temp_dir):
        for f in files:
            if f.lower().endswith(('.jpg', '.jpeg', '.png')):
                key = os.path.splitext(f)[0].lower().strip()
                image_map[key] = os.path.join(root, f)

    return image_map

# ================== MAIN ==================
def generate_documents(excel, word, zip_img, match_col, img_placeholder):
    temp_dir = tempfile.mkdtemp()

    try:
        # ===== EXCEL =====
        df = pd.read_excel(excel, dtype=str).fillna("")
        df.columns = df.columns.str.strip()

        if df.empty:
            st.error("Excel không có dữ liệu")
            return None

        # ===== WORD =====
        base_doc = Document(word)

        # ===== DEBUG PLACEHOLDER =====
        placeholders = extract_placeholders(base_doc)

        st.write("📌 Placeholder trong Word:")
        st.write([f"{{{{{p}}}}}" for p in placeholders])

        st.write("✅ Cột trong Excel:")
        st.write(df.columns.tolist())

        missing = [p for p in placeholders if p not in df.columns]
        extra = [c for c in df.columns if c not in placeholders]

        st.write("❌ Placeholder KHÔNG có trong Excel:", missing)
        st.write("⚠️ Cột Excel KHÔNG dùng:", extra)

        # ===== IMAGE ZIP =====
        img_map = {}
        if zip_img and match_col and img_placeholder:
            img_dir = os.path.join(temp_dir, "imgs")
            os.makedirs(img_dir, exist_ok=True)

            img_map = extract_images(zip_img, img_dir)

            df[img_placeholder] = df[match_col].apply(
                lambda x: img_map.get(str(x).lower().strip(), "")
            )

        # ===== TEMPLATE =====
        blocks = [copy.deepcopy(b) for b in base_doc.element.body]
        base_doc.element.body.clear()

        bar = st.progress(0)
        status = st.empty()

        for idx, row in df.iterrows():
            bar.progress((idx + 1) / len(df))
            status.text(f"Đang xử lý hồ sơ {idx+1}/{len(df)}")

            for b in blocks:
                base_doc.element.body.append(copy.deepcopy(b))

            # ===== PARAGRAPH =====
            for p in base_doc.paragraphs:
                for col in df.columns:
                    ph = f"{{{{{col}}}}}"
                    val = normalize_date(row[col])

                    if ph in p.text:
                        if col == img_placeholder and val:
                            insert_image(p, ph, val)
                        else:
                            replace_text_in_paragraph(p, ph, val)

            # ===== TABLE =====
            for table in base_doc.tables:
                for r in table.rows:
                    for c in r.cells:
                        for col in df.columns:
                            ph = f"{{{{{col}}}}}"
                            val = normalize_date(row[col])

                            if ph in c.text:
                                if col == img_placeholder and val:
                                    insert_image(c, ph, val)
                                else:
                                    replace_text_in_cell(c, ph, val)

            if idx < len(df) - 1:
                base_doc.add_page_break()

        out_path = os.path.join(temp_dir, "KetQua_DonDaDien.docx")
        base_doc.save(out_path)

        return out_path

    except Exception:
        st.error("Lỗi xử lý")
        st.code(traceback.format_exc())
        return None

# ================== UI ==================
st.title("📄 Tool Điền Đơn Tự Động")

c1, c2 = st.columns(2)

with c1:
    excel = st.file_uploader("1. Excel dữ liệu", type="xlsx")

with c2:
    word = st.file_uploader("2. Word mẫu", type="docx")

use_zip = st.checkbox("Có ảnh (.zip)")

zip_file = None
match_col = None
img_placeholder = ""

if use_zip and excel:
    zip_file = st.file_uploader("Zip ảnh", type="zip")
    cols = pd.read_excel(excel, nrows=0).columns.tolist()
    match_col = st.selectbox("Cột khớp ảnh", cols)
    img_placeholder = st.text_input("Placeholder ảnh", value="ANH_THE")

if st.button("🚀 BẮT ĐẦU"):
    if excel and word:
        with st.spinner("Đang xử lý..."):
            out = generate_documents(excel, word, zip_file, match_col, img_placeholder)

            if out:
                st.success("Hoàn tất!")
                with open(out, "rb") as f:
                    st.download_button(
                        "📥 Tải file kết quả",
                        f,
                        file_name="KetQua_DonDaDien.docx"
                    )
    else:
        st.warning("Thiếu Excel hoặc Word")
